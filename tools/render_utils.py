"""Pure rendering utilities for PPTX/DOCX generation. No crewai dependency."""

import os
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from docx.shared import Pt as DocPt, RGBColor as DocRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from tools.logging_config import get_logger

logger = get_logger("tools.render_utils")


def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip("#")
    if len(hex_str) == 3:
        hex_str = "".join([c * 2 for c in hex_str])
    return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


def hex_to_doc_rgb(hex_str):
    hex_str = hex_str.lstrip("#")
    if len(hex_str) == 3:
        hex_str = "".join([c * 2 for c in hex_str])
    return DocRGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


def chunk_bullets(bullets: list, max_per_slide: int = 6) -> list:
    """Split a list of bullets into chunks of max_per_slide each."""
    return [bullets[i:i + max_per_slide] for i in range(0, len(bullets), max_per_slide)]


def build_slide_entries(slides_data: list) -> list:
    """Expand slide data into individual slide entries, splitting overflow bullets."""
    slide_entries = []
    for data in slides_data:
        bullets = data.get("bullets", [])
        chunks = chunk_bullets(bullets)
        for idx, chunk in enumerate(chunks):
            title = data.get("title", "Untitled")
            if len(chunks) > 1:
                title += f" ({idx+1}/{len(chunks)})"
            slide_entries.append({
                "title": title,
                "bullets": chunk,
                "diagram_path": data.get("diagram_path") if idx == 0 else None,
                "image_path": data.get("image_path") if idx == 0 else None,
                "layout_type": data.get("layout_type", "Title and Content"),
                "subtitle": data.get("subtitle") if idx == 0 else None,
                "attribution": data.get("attribution") if idx == 0 else None,
            })
    return slide_entries


def resolve_image_path(diagram_path, image_path):
    candidates = []
    if diagram_path:
        candidates.append(diagram_path)
        candidates.append(os.path.join("temp", diagram_path))
    if image_path:
        candidates.append(image_path)
        candidates.append(os.path.join("temp", image_path))
    for p in candidates:
        if p and os.path.exists(p):
            return os.path.abspath(p)
    return None


def add_accent_bar(slide, accent_rgb, top, height):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, top, Inches(0.06), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_rgb
    bar.line.visible = False


def apply_enhanced_ppt_style(slide, title_text, bullets, style, diagram_path=None, image_path=None, slide_number=None):
    primary_rgb = hex_to_rgb(style.get("primary_color", "#0F2C59"))
    secondary_rgb = hex_to_rgb(style.get("secondary_color", "#1E5F74"))
    accent_rgb = hex_to_rgb(style.get("accent_color", "#E8B931"))
    bg_rgb = hex_to_rgb(style.get("background_color", "#F5F7FA"))
    text_rgb = hex_to_rgb(style.get("text_primary", "#1A1A2E"))
    inv_rgb = hex_to_rgb(style.get("text_inverse", "#FFFFFF"))
    header_size = style.get("header_font_size", 36)
    body_size = style.get("body_font_size", 20)
    font_family = style.get("font_family", "Calibri")

    # Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_rgb

    # Header bar
    header_h = Inches(1.2)
    header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), header_h)
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = primary_rgb
    header_box.line.visible = False

    # Title
    title_box = slide.shapes.add_textbox(0, 0, Inches(13.33), header_h)
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(header_size)
    p.font.bold = True
    p.font.name = font_family
    p.font.color.rgb = inv_rgb

    # Content area
    left_margin = Inches(0.5)
    top_margin = Inches(1.5)
    content_height = Inches(5.3)

    final_image_path = resolve_image_path(diagram_path, image_path)

    if final_image_path:
        text_width = Inches(6.0)
        img_left = Inches(6.8)
        img_max_width = Inches(6.0)
        img_max_height = Inches(5.3)

        text_box = slide.shapes.add_textbox(left_margin, top_margin, text_width, content_height)
        add_accent_bar(slide, accent_rgb, top_margin, content_height)

        try:
            from PIL import Image as PILImage
            with PILImage.open(final_image_path) as img:
                orig_w, orig_h = img.size
            aspect = orig_h / orig_w if orig_w else 1
            target_w = img_max_width
            target_h = target_w * aspect
            if target_h > img_max_height:
                target_h = img_max_height
                target_w = target_h / aspect if aspect else img_max_width
            img_top = top_margin + (content_height - target_h) / 2
            slide.shapes.add_picture(final_image_path, img_left, img_top, width=target_w, height=target_h)
        except Exception:
            try:
                slide.shapes.add_picture(final_image_path, img_left, top_margin, width=img_max_width)
            except Exception as e:
                logger.warning("Could not add image %s: %s", final_image_path, e)
    else:
        text_width = Inches(12.33)
        text_box = slide.shapes.add_textbox(left_margin, top_margin, text_width, content_height)
        add_accent_bar(slide, accent_rgb, top_margin, content_height)

    # Bullet text
    tf = text_box.text_frame
    tf.word_wrap = True
    for i, point in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        text = point.text if hasattr(point, "text") else str(point)
        level = point.level if hasattr(point, "level") else 0
        p.text = f"• {text}"
        p.font.size = Pt(body_size)
        p.font.name = font_family
        p.font.color.rgb = text_rgb
        p.level = max(0, min(int(level), 3))
        p.space_before = Pt(10)
        p.space_after = Pt(6)
        p.alignment = PP_ALIGN.LEFT

    # Footer bar
    footer_h = Inches(0.3)
    footer_y = Inches(7.2)
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, footer_y, Inches(13.33), footer_h)
    footer.fill.solid()
    footer.fill.fore_color.rgb = primary_rgb
    footer.line.visible = False

    # Slide number
    if slide_number:
        num_box = slide.shapes.add_textbox(Inches(12.0), footer_y, Inches(1.0), footer_h)
        ntf = num_box.text_frame
        ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
        np = ntf.paragraphs[0]
        np.text = str(slide_number)
        np.alignment = PP_ALIGN.RIGHT
        np.font.size = Pt(12)
        np.font.name = font_family
        np.font.color.rgb = inv_rgb


def render_title_slide(slide, title_text, style, subtitle=None, slide_number=None):
    """Render a title slide with large centered title and optional subtitle."""
    primary_rgb = hex_to_rgb(style.get("primary_color", "#0F2C59"))
    bg_rgb = hex_to_rgb(style.get("background_color", "#F5F7FA"))
    text_rgb = hex_to_rgb(style.get("text_primary", "#1A1A2E"))
    inv_rgb = hex_to_rgb(style.get("text_inverse", "#FFFFFF"))
    font_family = style.get("font_family", "Calibri")

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_rgb

    # Full-bleed header bar
    header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), Inches(3.0))
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = primary_rgb
    header_box.line.visible = False

    # Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.33), Inches(2.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.name = font_family
    p.font.color.rgb = inv_rgb

    # Subtitle
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.33), Inches(0.8))
        stf = sub_box.text_frame
        stf.word_wrap = True
        sp = stf.paragraphs[0]
        sp.text = subtitle
        sp.alignment = PP_ALIGN.CENTER
        sp.font.size = Pt(24)
        sp.font.name = font_family
        sp.font.color.rgb = inv_rgb

    # Date in footer area
    from datetime import datetime
    date_box = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11.33), Inches(0.5))
    dtf = date_box.text_frame
    dp = dtf.paragraphs[0]
    dp.text = datetime.now().strftime("%B %d, %Y")
    dp.alignment = PP_ALIGN.CENTER
    dp.font.size = Pt(14)
    dp.font.name = font_family
    dp.font.color.rgb = text_rgb


def render_agenda_slide(slide, title_text, agenda_items, style, slide_number=None):
    """Render an agenda slide with a numbered list of upcoming slides."""
    primary_rgb = hex_to_rgb(style.get("primary_color", "#0F2C59"))
    secondary_rgb = hex_to_rgb(style.get("secondary_color", "#1E5F74"))
    bg_rgb = hex_to_rgb(style.get("background_color", "#F5F7FA"))
    text_rgb = hex_to_rgb(style.get("text_primary", "#1A1A2E"))
    inv_rgb = hex_to_rgb(style.get("text_inverse", "#FFFFFF"))
    font_family = style.get("font_family", "Calibri")

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_rgb

    # Header bar
    header_h = Inches(1.2)
    header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), header_h)
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = primary_rgb
    header_box.line.visible = False

    title_box = slide.shapes.add_textbox(0, 0, Inches(13.33), header_h)
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = font_family
    p.font.color.rgb = inv_rgb

    # Two-column agenda items
    left_margin = Inches(1.0)
    top_margin = Inches(1.8)
    col_width = Inches(5.5)
    mid_col = Inches(6.8)

    for i, item in enumerate(agenda_items):
        text = item.text if hasattr(item, "text") else str(item)
        col = i % 2
        row = i // 2
        x = left_margin if col == 0 else mid_col
        y = top_margin + Inches(0.5) * row
        item_box = slide.shapes.add_textbox(x, y, col_width, Inches(0.45))
        itf = item_box.text_frame
        ip = itf.paragraphs[0]
        ip.text = f"{i + 1}. {text}"
        ip.font.size = Pt(18)
        ip.font.name = font_family
        ip.font.color.rgb = text_rgb


def render_diagram_full_slide(slide, title_text, style, diagram_path=None, image_path=None, caption=None, slide_number=None):
    """Render a slide dominated by a single large diagram with optional caption."""
    primary_rgb = hex_to_rgb(style.get("primary_color", "#0F2C59"))
    bg_rgb = hex_to_rgb(style.get("background_color", "#F5F7FA"))
    inv_rgb = hex_to_rgb(style.get("text_inverse", "#FFFFFF"))
    font_family = style.get("font_family", "Calibri")

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_rgb

    # Small header
    header_h = Inches(0.9)
    header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), header_h)
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = primary_rgb
    header_box.line.visible = False

    title_box = slide.shapes.add_textbox(0, 0, Inches(13.33), header_h)
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.name = font_family
    p.font.color.rgb = inv_rgb

    # Large image area
    final_image_path = resolve_image_path(diagram_path, image_path)
    if final_image_path:
        img_left = Inches(1.5)
        img_top = Inches(1.2)
        img_max_width = Inches(10.33)
        img_max_height = Inches(5.5)
        try:
            from PIL import Image as PILImage
            with PILImage.open(final_image_path) as img:
                orig_w, orig_h = img.size
            aspect = orig_h / orig_w if orig_w else 1
            target_w = img_max_width
            target_h = target_w * aspect
            if target_h > img_max_height:
                target_h = img_max_height
                target_w = target_h / aspect if aspect else img_max_width
            img_top_centered = img_top + (Inches(5.5) - target_h) / 2
            slide.shapes.add_picture(final_image_path, img_left, img_top_centered, width=target_w, height=target_h)
        except Exception as e:
            logger.warning("Could not add diagram image: %s", e)

    # Caption
    if caption:
        cap_box = slide.shapes.add_textbox(Inches(1.5), Inches(6.8), Inches(10.33), Inches(0.5))
        ctf = cap_box.text_frame
        cp = ctf.paragraphs[0]
        cp.text = caption
        cp.alignment = PP_ALIGN.CENTER
        cp.font.size = Pt(12)
        cp.font.name = font_family
        cp.font.color.rgb = hex_to_rgb(style.get("secondary_color", "#1E5F74"))


def render_quote_slide(slide, title_text, style, quote_text=None, attribution=None, slide_number=None):
    """Render a quote/callout slide with large pull-quote text."""
    primary_rgb = hex_to_rgb(style.get("primary_color", "#0F2C59"))
    accent_rgb = hex_to_rgb(style.get("accent_color", "#E8B931"))
    bg_rgb = hex_to_rgb(style.get("background_color", "#F5F7FA"))
    text_rgb = hex_to_rgb(style.get("text_primary", "#1A1A2E"))
    font_family = style.get("font_family", "Calibri")

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_rgb

    # Accent bar on left
    add_accent_bar(slide, accent_rgb, Inches(1.0), Inches(5.5))

    # Quote text
    if quote_text:
        quote_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.33), Inches(4.5))
        qtf = quote_box.text_frame
        qtf.word_wrap = True
        qp = qtf.paragraphs[0]
        qp.text = f'"{quote_text}"'
        qp.font.size = Pt(28)
        qp.font.italic = True
        qp.font.name = font_family
        qp.font.color.rgb = text_rgb

    # Attribution
    if attribution:
        attr_box = slide.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(11.33), Inches(0.5))
        atf = attr_box.text_frame
        ap = atf.paragraphs[0]
        ap.text = f"— {attribution}"
        ap.font.size = Pt(18)
        ap.font.name = font_family
        ap.font.color.rgb = primary_rgb


def render_comparison_slide(slide, title_text, bullets, style, slide_number=None):
    """Render a two-column comparison slide."""
    primary_rgb = hex_to_rgb(style.get("primary_color", "#0F2C59"))
    bg_rgb = hex_to_rgb(style.get("background_color", "#F5F7FA"))
    text_rgb = hex_to_rgb(style.get("text_primary", "#1A1A2E"))
    inv_rgb = hex_to_rgb(style.get("text_inverse", "#FFFFFF"))
    font_family = style.get("font_family", "Calibri")

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_rgb

    # Header
    header_h = Inches(1.2)
    header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.33), header_h)
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = primary_rgb
    header_box.line.visible = False

    title_box = slide.shapes.add_textbox(0, 0, Inches(13.33), header_h)
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = font_family
    p.font.color.rgb = inv_rgb

    # Two columns
    left_margin = Inches(0.8)
    col_width = Inches(5.8)
    mid = Inches(6.8)
    top_margin = Inches(1.6)

    mid_idx = len(bullets) // 2 if bullets else 0
    for col, start, end in [(0, 0, mid_idx), (1, mid_idx, len(bullets))]:
        x = left_margin if col == 0 else mid
        col_box = slide.shapes.add_textbox(x, top_margin, col_width, Inches(5.0))
        ctf = col_box.text_frame
        ctf.word_wrap = True
        for i, point in enumerate(bullets[start:end]):
            text = point.text if hasattr(point, "text") else str(point)
            if i == 0:
                pp = ctf.paragraphs[0]
            else:
                pp = ctf.add_paragraph()
            pp.text = f"• {text}"
            pp.font.size = Pt(18)
            pp.font.name = font_family
            pp.font.color.rgb = text_rgb
            pp.space_before = Pt(8)


def render_divider_slide(slide, title_text, style, slide_number=None):
    """Render a section divider slide with large centered section title."""
    primary_rgb = hex_to_rgb(style.get("primary_color", "#0F2C59"))
    accent_rgb = hex_to_rgb(style.get("accent_color", "#E8B931"))
    bg_rgb = hex_to_rgb(style.get("background_color", "#F5F7FA"))
    inv_rgb = hex_to_rgb(style.get("text_inverse", "#FFFFFF"))
    font_family = style.get("font_family", "Calibri")

    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_rgb

    # Accent bar top
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(3.0), Inches(13.33), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_rgb
    bar.line.visible = False

    # Centered title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(3.2), Inches(11.33), Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.name = font_family
    p.font.color.rgb = primary_rgb


def dispatch_slide_layout(slide, entry, style, slide_number=None):
    """Route a slide entry to the appropriate layout renderer based on layout_type."""
    layout = entry.get("layout_type", "Title and Content")
    title = entry.get("title", "Untitled")
    bullets = entry.get("bullets", [])
    diagram_path = entry.get("diagram_path")
    image_path = entry.get("image_path")

    if layout == "Title":
        render_title_slide(slide, title, style, subtitle=entry.get("subtitle"), slide_number=slide_number)
    elif layout == "Agenda":
        render_agenda_slide(slide, title, bullets, style, slide_number=slide_number)
    elif layout == "Diagram Full":
        render_diagram_full_slide(slide, title, style, diagram_path=diagram_path, image_path=image_path, slide_number=slide_number)
    elif layout == "Quote":
        first = bullets[0] if bullets else None
        if first:
            quote_text = first.get("text") if isinstance(first, dict) else getattr(first, "text", str(first))
        else:
            quote_text = None
        render_quote_slide(slide, title, style, quote_text=quote_text, attribution=entry.get("attribution"), slide_number=slide_number)
    elif layout == "Comparison":
        render_comparison_slide(slide, title, bullets, style, slide_number=slide_number)
    elif layout == "Divider":
        render_divider_slide(slide, title, style, slide_number=slide_number)
    else:
        apply_enhanced_ppt_style(slide, title, bullets, style, diagram_path=diagram_path, image_path=image_path, slide_number=slide_number)


def generate_preview_html(run_id, slide_entries, output_dir="output"):
    """Generate an HTML preview file with a grid of slide titles.

    Creates `output/preview_{run_id}.html` showing all slide titles in a
    numbered grid for quick visual verification without opening the PPT.

    Args:
        run_id: The run identifier for filename.
        slide_entries: List of slide entry dicts (from build_slide_entries).
        output_dir: Directory to write the HTML file.

    Returns:
        Absolute path to the generated HTML file.
    """
    os.makedirs(output_dir, exist_ok=True)
    lines = [
        "<!DOCTYPE html>",
        "<html><head><meta charset='utf-8'>",
        f"<title>Preview — {run_id}</title>",
        "<style>",
        "body{font-family:Calibri,Arial,sans-serif;background:#f5f7fa;margin:0;padding:20px;}",
        "h1{color:#0F2C59;}",
        ".grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(250px,1fr));gap:16px;margin-top:20px;}",
        ".slide-card{background:#fff;border-radius:8px;padding:16px;box-shadow:0 1px 3px rgba(0,0,0,0.12);border-left:4px solid #0F2C59;}",
        ".slide-num{font-size:12px;color:#666;margin-bottom:4px;}",
        ".slide-title{font-size:16px;font-weight:600;color:#1A1A2E;}",
        ".slide-count{margin-top:20px;color:#666;font-size:14px;}",
        "</style></head><body>",
        f"<h1>Presentation Preview — {run_id}</h1>",
        f"<div class='slide-count'>{len(slide_entries)} slides</div>",
        "<div class='grid'>",
    ]

    for i, entry in enumerate(slide_entries, start=1):
        title = entry.get("title", "Untitled")
        layout = entry.get("layout_type", "Title and Content")
        lines.append(f"<div class='slide-card'>")
        lines.append(f"<div class='slide-num'>Slide {i} — {layout}</div>")
        lines.append(f"<div class='slide-title'>{title}</div>")
        lines.append(f"</div>")

    lines.append("</div></body></html>")

    path = os.path.join(output_dir, f"preview_{run_id}.html")
    with open(path, "w") as f:
        f.write("\n".join(lines))

    logger.info("Preview HTML saved: %s", path)
    return os.path.abspath(path)


def apply_reference_doc_style(doc, sections, style_config):
    """Apply themed styling to a reference manual document with conditional subheadings."""
    primary_color = style_config.get("primary_color", "#0F2C59")
    secondary_color = style_config.get("secondary_color", "#1E5F74")
    text_primary = style_config.get("text_primary", "#1A1A2E")

    title = doc.add_heading("Technical Reference Manual", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = hex_to_doc_rgb(primary_color)
    title.runs[0].font.size = DocPt(22)

    for section in sections:
        heading = section.get("heading", "Section") if isinstance(section, dict) else getattr(section, "heading", "Section")
        h1 = doc.add_heading(heading, level=1)
        h1.runs[0].font.color.rgb = hex_to_doc_rgb(primary_color)
        h1.runs[0].font.size = DocPt(18)

        def _add_doc_heading(text, level=2):
            h = doc.add_heading(text, level=level)
            h.runs[0].font.color.rgb = hex_to_doc_rgb(secondary_color)
            h.runs[0].font.size = DocPt(14)
            return h

        def _add_doc_paragraph(text):
            p = doc.add_paragraph(text)
            p.paragraph_format.space_after = DocPt(12)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                run.font.color.rgb = hex_to_doc_rgb(text_primary)
            return p

        # Custom subheadings or defaults
        if isinstance(section, dict):
            prose = section.get("prose_content", "")
            impl = section.get("implementation_guide", "")
            trouble = section.get("troubleshooting")
            custom = section.get("custom_subheadings")
        else:
            prose = section.prose_content
            impl = section.implementation_guide
            trouble = section.troubleshooting
            custom = getattr(section, "custom_subheadings", None)

        if custom:
            for sh in custom:
                title_text = sh.get("title", "") if isinstance(sh, dict) else getattr(sh, "title", "")
                level = sh.get("level", 2) if isinstance(sh, dict) else getattr(sh, "level", 2)
                _add_doc_heading(title_text, level=level)
        else:
            _add_doc_heading("Technical Analysis", level=2)

        _add_doc_paragraph(prose)
        _add_doc_heading("Implementation Details", level=2)
        _add_doc_paragraph(impl)

        if trouble:
            _add_doc_heading("Troubleshooting & Security Considerations", level=2)
            _add_doc_paragraph(trouble)

        doc.add_page_break()
